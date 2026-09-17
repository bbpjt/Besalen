import os
import re
import json
import math
import zipfile
import openpyxl
import docx

KAB_COORDS = {
    "Kabupaten Cilacap": (-7.545, 108.995),
    "Kabupaten Banyumas": (-7.455, 109.245),
    "Kabupaten Purbalingga": (-7.345, 109.355),
    "Kabupaten Banjarnegara": (-7.385, 109.685),
    "Kabupaten Kebumen": (-7.655, 109.665),
    "Kabupaten Purworejo": (-7.715, 110.015),
    "Kabupaten Wonosobo": (-7.365, 109.915),
    "Kabupaten Magelang": (-7.535, 110.215),
    "Kabupaten Boyolali": (-7.525, 110.605),
    "Kabupaten Klaten": (-7.705, 110.605),
    "Kabupaten Sukoharjo": (-7.675, 110.835),
    "Kabupaten Wonogiri": (-7.825, 110.925),
    "Kabupaten Karanganyar": (-7.605, 110.955),
    "Kabupaten Sragen": (-7.425, 111.025),
    "Kabupaten Grobogan": (-7.115, 110.915),
    "Kabupaten Blora": (-7.015, 111.415),
    "Kabupaten Rembang": (-6.775, 111.345),
    "Kabupaten Pati": (-6.755, 111.035),
    "Kabupaten Kudus": (-6.805, 110.845),
    "Kabupaten Jepara": (-6.585, 110.675),
    "Kabupaten Demak": (-6.895, 110.645),
    "Kabupaten Semarang": (-7.215, 110.425),
    "Kabupaten Temanggung": (-7.275, 110.155),
    "Kabupaten Kendal": (-6.995, 110.185),
    "Kabupaten Batang": (-7.025, 109.845),
    "Kabupaten Pekalongan": (-7.055, 109.625),
    "Kabupaten Pemalang": (-7.015, 109.385),
    "Kabupaten Tegal": (-7.025, 109.155),
    "Kabupaten Brebes": (-7.015, 108.945),
    "Kota Magelang": (-7.475, 110.220),
    "Kota Surakarta": (-7.570, 110.825),
    "Kota Salatiga": (-7.330, 110.505),
    "Kota Semarang": (-6.995, 110.420),
    "Kota Pekalongan": (-6.890, 109.675),
    "Kota Tegal": (-6.870, 109.135)
}

def get_centroid(kab_name):
    clean_name = kab_name.strip()
    for k, v in KAB_COORDS.items():
        if k.lower() in clean_name.lower() or clean_name.lower() in k.lower():
            return v
    return (-7.15, 110.14)

def apply_jitter(base_lat, base_lng, item_index, seed_offset=0):
    angle = (item_index * 137.508 + seed_offset) * (math.pi / 180.0)
    radius = 0.015 + (item_index % 4) * 0.010
    jitter_lat = round(base_lat + radius * math.sin(angle), 6)
    jitter_lng = round(base_lng + (radius / math.cos(math.radians(base_lat))) * math.cos(angle), 6)
    return jitter_lat, jitter_lng

def load_verified_sources():
    sources_map = {}
    if not os.path.exists("Sumber_Ilmiah_Verifikasi_Sastra_Lisan_Jateng.xlsx"):
        return sources_map
    wb = openpyxl.load_workbook("Sumber_Ilmiah_Verifikasi_Sastra_Lisan_Jateng.xlsx", data_only=True)
    s = wb["Verifikasi Sumber"]
    for r in range(6, s.max_row + 1):
        nama = s.cell(r, 2).value
        if not nama: continue
        nama = str(nama).strip()
        src1 = str(s.cell(r, 7).value or "").strip()
        url1 = str(s.cell(r, 8).value or "").strip()
        src2 = str(s.cell(r, 10).value or "").strip()
        url2 = str(s.cell(r, 11).value or "").strip()
        
        # Clean URL if it's "None" or text
        if not url1.startswith("http"): url1 = ""
        if not url2.startswith("http"): url2 = ""
        
        sources_map[nama.lower()] = {
            "sumber1": src1,
            "url1": url1,
            "sumber2": src2,
            "url2": url2
        }
    return sources_map

def get_kentrung_photos():
    photos = []
    if os.path.exists("media/kentrung"):
        for f in sorted(os.listdir("media/kentrung")):
            if any(f.lower().endswith(ext) for ext in [".jpg", ".jpeg", ".png"]):
                photos.append(f"media/kentrung/{f}")
    if not photos and os.path.exists("kentrung-Cover.jpg"):
        photos.append("kentrung-Cover.jpg")
    return photos

def main():
    print("Re-extracting dataset with rich source URLs & Kentrung Drive photos...")
    sources_map = load_verified_sources()
    kentrung_photos = get_kentrung_photos()
    
    banyumas_photos = [f"media/banyumas/{f}" for f in sorted(os.listdir("media/banyumas"))] if os.path.exists("media/banyumas") else []
    wonosobo_photos = [f"media/wonosobo/{f}" for f in sorted(os.listdir("media/wonosobo"))] if os.path.exists("media/wonosobo") else []
    
    excel_path = "Peta_Sastra_Lisan_Jawa_Tengah_Fondasi_Terverifikasi.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. Ring 1
    s1 = wb["Ring 1 - Tradisi Emas"]
    ring1 = []
    for r in range(2, 5):
        row = [s1.cell(r, c).value for c in range(1, 15)]
        reg_id, nama, kab, kec, desa, maestro, usia, kom, bentuk, musik, rep = row[:11]
        
        if "Blora" in nama:
            lat, lng = -6.9944, 111.3789
            karesidenan = "Karesidenan Pati (Muria Raya)"
            zona = "Karst Kendeng & Hutan Jati Rimba"
            yt_id = "5XteEv2MU_g"
            yt_url = "https://youtu.be/5XteEv2MU_g"
            local_vid = "kentrung.mp4"
            cover = "kentrung-Cover.jpg"
            gallery = kentrung_photos
            ringkasan = "Kentrung Blora adalah seni tutur lisan berirama bernuansa Islam-Jawa pesisir dengan tabuhan terbang besar dan ketunthung. Penutur membawakan carita kepahlawanan dan religius dengan variasi kandha, janturan, dan caturan secara soliter."
            sumber = "Balai Bahasa Provinsi Jawa Tengah (2026); Transkripsi Fonemik 367 Segmen"
            url1 = "https://youtu.be/5XteEv2MU_g"
            url2 = "https://dapobud.kemenbud.go.id/wbtb"
        elif "Babad Pasir Luhur" in nama:
            lat, lng = -7.3889, 109.1867
            karesidenan = "Karesidenan Banyumas"
            zona = "Lembah DAS Serayu & Lereng Gunung Slamet"
            yt_id = "Ixm0NVVzniM"
            yt_url = "https://youtu.be/Ixm0NVVzniM"
            local_vid = ""
            cover = banyumas_photos[0] if banyumas_photos else ""
            gallery = banyumas_photos[:18]
            ringkasan = "Maca Babad Pasir Luhur merupakan tradisi pelantunan tembang macapat (sekar macapat) dari naskah babad tertua di lembah Serayu lereng Gunung Cokol. Pertunjukan memadukan tradisi sastra tulis babad dengan sastra tutur jemblungan akapela."
            sumber = "Balai Bahasa Provinsi Jawa Tengah (2026); Laporan Revitalisasi Maca Babad Pasir Luhur"
            url1 = "https://youtu.be/Ixm0NVVzniM"
            url2 = "https://dapobud.kemenbud.go.id/wbtb"
        else:
            lat, lng = -7.4206, 109.8433
            karesidenan = "Karesidenan Kedu"
            zona = "Dataran Tinggi Serayu Hulu & Pegunungan Dieng"
            yt_id = "Q_JAmcKSFB0"
            yt_url = "https://youtube.com/live/Q_JAmcKSFB0"
            local_vid = ""
            cover = wonosobo_photos[0] if wonosobo_photos else ""
            gallery = wonosobo_photos
            ringkasan = "Wayang Othok Obrol adalah teater tutur wayang purwa komunal yang sangat langka di Selokromo Leksono Wonosobo. Menggunakan wayang kulit tua peninggalan leluhur, dalang menuturkan pakeliran, suluk, dan dialog tanpa gamelan lengkap, melainkan iringan kendang dan senggakan mulut."
            sumber = "Balai Bahasa Provinsi Jawa Tengah (2026); Laporan Revitalisasi Wayang Othok Obrol"
            url1 = "https://youtube.com/live/Q_JAmcKSFB0"
            url2 = "https://dapobud.kemenbud.go.id/wbtb"
            
        ring1.append({
            "id": reg_id,
            "nama": nama,
            "ring": "Ring 1 - Terverifikasi",
            "ring_level": 1,
            "status_label": "TERVERIFIKASI (VALIDASI LAPANGAN PENUH)",
            "kabupaten": kab,
            "kecamatan": kec,
            "desa": desa,
            "karesidenan": karesidenan,
            "zona_ekologi": zona,
            "maestro": maestro,
            "usia_garis": usia,
            "komunitas": kom,
            "bentuk_tuturan": bentuk,
            "iringan_musik": musik,
            "repertoar": rep,
            "latitude": lat,
            "longitude": lng,
            "youtube_id": yt_id,
            "youtube_url": yt_url,
            "video_local": local_vid,
            "cover_image": cover,
            "gallery_images": gallery,
            "ringkasan_ilmiah": ringkasan,
            "sumber_ilmiah": sumber,
            "url1": url1,
            "url2": url2
        })
        
    # 2. Ring 2
    s2 = wb["Ring 2 - Tervalidasi Teks (44)"]
    ring2 = []
    for r in range(2, 46):
        row = [s2.cell(r, c).value for c in range(1, 14)]
        no, nama, kab, karesidenan, zona, stat, unsur, ringkasan, src1, dasar1, src2 = row[:11]
        kab = kab or "Jawa Tengah"
        base_lat, base_lng = get_centroid(kab)
        lat, lng = apply_jitter(base_lat, base_lng, r, seed_offset=42)
        
        # Link from sources_map
        mapped = sources_map.get(str(nama).strip().lower(), {})
        url1 = mapped.get("url1", "")
        url2 = mapped.get("url2", "")
        
        # Extract direct url if embedded inside src1 text
        if not url1 and "http" in str(src1):
            m = re.search(r'(https?://[^\s]+)', str(src1))
            if m: url1 = m.group(1)
            
        ring2.append({
            "id": f"SLJT-R2-{int(no):03d}",
            "nama": nama,
            "ring": "Ring 2 - Ada referensi",
            "ring_level": 2,
            "status_label": "TERVERIFIKASI TEKS ILMIAH",
            "kabupaten": kab,
            "karesidenan": karesidenan or "Jawa Tengah",
            "zona_ekologi": zona or "Kawasan Budaya Jawa",
            "status_verifikasi": "TERVERIFIKASI TEKS (BERTEKS)",
            "unsur_teks": unsur or "",
            "ringkasan_ilmiah": ringkasan or "",
            "sumber_ilmiah_1": src1 or "",
            "dasar_bukti_1": dasar1 or "",
            "sumber_ilmiah_2": src2 or "",
            "url1": url1,
            "url2": url2,
            "latitude": lat,
            "longitude": lng
        })
        
    # 3. Ring 3
    s3 = wb["Ring 3 - Tercatat WBTB (44)"]
    ring3 = []
    for r in range(2, 46):
        row = [s3.cell(r, c).value for c in range(1, 12)]
        no, nama, kab, karesidenan, zona, stat, unsur, catatan, src = row[:9]
        kab = kab or "Jawa Tengah"
        base_lat, base_lng = get_centroid(kab)
        lat, lng = apply_jitter(base_lat, base_lng, r, seed_offset=199)
        
        mapped = sources_map.get(str(nama).strip().lower(), {})
        url1 = mapped.get("url1", "")
        if not url1 and "http" in str(src):
            m = re.search(r'(https?://[^\s]+)', str(src))
            if m: url1 = m.group(1)
            
        ring3.append({
            "id": f"SLJT-R3-{int(no):03d}",
            "nama": nama,
            "ring": "Ring 3 - Tercatat WBTB",
            "ring_level": 3,
            "status_label": "PERLU VERIFIKASI LAPANGAN",
            "kabupaten": kab,
            "karesidenan": karesidenan or "Jawa Tengah",
            "zona_ekologi": zona or "Kawasan Budaya Jawa",
            "status_kerja": "PERLU VERIFIKASI (PRIORITAS OBSERVASI)",
            "unsur_teks": unsur or "",
            "catatan_kritis": catatan or "",
            "sumber_referensi": src or "",
            "url1": url1,
            "latitude": lat,
            "longitude": lng
        })
        
    # 4. Ring 4
    s4 = wb["Ring 4 - Eksklusi Non-Sastra"]
    ring4 = []
    for r in range(2, s4.max_row + 1):
        row = [s4.cell(r, c).value for c in range(1, 10)]
        no = row[0]
        if no is None: continue
        nama, kab = row[1], row[2] or "Jawa Tengah"
        kategori_asli = row[4] or ""
        alasan = row[5] or ""
        sumber = row[6] or ""
        
        mapped = sources_map.get(str(nama).strip().lower(), {})
        url1 = mapped.get("url1", "")
        if not url1 and "http" in str(sumber):
            m = re.search(r'(https?://[^\s]+)', str(sumber))
            if m: url1 = m.group(1)
            
        ring4.append({
            "no": int(no),
            "nama": nama,
            "kabupaten": kab,
            "status_keputusan": "DIEKSKLUSI RESMI",
            "kategori_asli": kategori_asli,
            "alasan_eksklusi": alasan,
            "sumber": sumber,
            "url1": url1
        })
        
    # 5. Peta 35 Kab-Kota
    s5 = wb["Peta 35 Kab-Kota"]
    kabupaten_list = []
    for r in range(2, 37):
        row = [s5.cell(r, c).value for c in range(1, 13)]
        kode = str(row[0]).strip()
        nama_kab = str(row[1]).strip()
        lat, lng = get_centroid(nama_kab)
        kabupaten_list.append({
            "kode_wilayah": kode,
            "nama": nama_kab,
            "karesidenan": str(row[2]).strip(),
            "zona_ekologi": str(row[3]).strip(),
            "r1_count": int(row[4] or 0),
            "r2_count": int(row[5] or 0),
            "r3_count": int(row[6] or 0),
            "r4_count": int(row[7] or 0),
            "total_potensi": int(row[4] or 0) + int(row[5] or 0) + int(row[6] or 0),
            "status_keterwakilan": str(row[9] or "").strip(),
            "rencana_tindak_lanjut": str(row[10] or "").strip(),
            "latitude": lat,
            "longitude": lng
        })
        
    master_data = {
        "metadata": {
            "judul": "Peta Sastra Lisan di Jawa Tengah",
            "tahun": 2026,
            "instansi": "Balai Bahasa Provinsi Jawa Tengah",
            "total_ring1": len(ring1),
            "total_ring2": len(ring2),
            "total_ring3": len(ring3),
            "total_ring4": len(ring4),
            "total_kabupaten": len(kabupaten_list)
        },
        "ring1": ring1,
        "ring2": ring2,
        "ring3": ring3,
        "ring4": ring4,
        "kabupaten": kabupaten_list
    }
    
    with open("data/sastra_data.js", "w", encoding="utf-8") as f:
        f.write("/**\n * Peta Sastra Lisan di Jawa Tengah\n * Balai Bahasa Provinsi Jawa Tengah\n */\n")
        f.write("window.SASTRA_DATA = ")
        f.write(json.dumps(master_data, indent=2, ensure_ascii=False))
        f.write(";\n")
        
    print(f"Updated data/sastra_data.js with Kentrung photos ({len(kentrung_photos)}) and clickable source URLs!")

if __name__ == "__main__":
    main()
